"""模板占位符扫描（模板自适应方案 3.1）：提取与分类的唯一实现。

提取规则：docx 段落 + 表格单元格中的 {{占位符}}（与测试漂移检测同一正则）。
分类启发式（输出建议清单，人工确认后落盘——不承诺全自动）：
- 表格单元格内 / 嵌在句子中间           → 标量（确信）
- 独立段落且以"表格"结尾               → 动态表候选（确信：哨兵段落插真表）
- 其余独立段落                         → 待判别（文本段 or 独立标量如"报告标题"，
                                         需人工或对照现行注册表归属）
"""
from __future__ import annotations

import re
from pathlib import Path

import docx

_PLACEHOLDER_RE = re.compile(r"\{\{([^}]+)\}\}")


def _is_standalone(paragraph_text: str, placeholder: str) -> bool:
    """占位符是否独占整个段落（哨兵段落特征）。"""
    return paragraph_text.strip() == f"{{{{{placeholder}}}}}"


def extract_placeholders(template: Path) -> dict:
    """返回 {"all", "scalars", "dynamic_tables", "standalone_other"}：
    - scalars / dynamic_tables 为确信分类；
    - standalone_other 为独立段落中的非"表格"键（文本段与独立标量无法靠位置区分，
      落盘前需人工确认——"报告标题"事件：独立段落 ≠ 一定是文本段）。"""
    d = docx.Document(str(template))
    embedded: set[str] = set()
    standalone: set[str] = set()
    in_tables: set[str] = set()

    for p in d.paragraphs:
        for ph in _PLACEHOLDER_RE.findall(p.text):
            (standalone if _is_standalone(p.text, ph) else embedded).add(ph)
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                for ph in _PLACEHOLDER_RE.findall(c.text):
                    in_tables.add(ph)

    dynamic_tables = {ph for ph in standalone if ph.endswith("表格")}
    standalone_other = standalone - dynamic_tables
    scalars = embedded | in_tables
    return {"all": scalars | dynamic_tables | standalone_other,
            "scalars": scalars, "dynamic_tables": dynamic_tables,
            "standalone_other": standalone_other}


def diff_against_registry(found: dict) -> dict:
    """与现行注册表比对：{"missing": 模板有而注册表无, "stale": 注册表有而模板无,
    "misclassified": 确信分类与注册表冲突的键}；三者皆空 = 无漂移。
    standalone_other 键按现行注册表归属判定，不算误分类（待人工判别桶）。"""
    from .report_registry import (DYNAMIC_TABLE_PLACEHOLDERS,
                                  SCALAR_PLACEHOLDERS,
                                  TEXT_SECTION_PLACEHOLDERS)
    reg_all = SCALAR_PLACEHOLDERS | DYNAMIC_TABLE_PLACEHOLDERS | TEXT_SECTION_PLACEHOLDERS
    misclassified = (
        {k for k in found["dynamic_tables"] & (SCALAR_PLACEHOLDERS | TEXT_SECTION_PLACEHOLDERS)}
        | {k for k in found["scalars"] & (DYNAMIC_TABLE_PLACEHOLDERS | TEXT_SECTION_PLACEHOLDERS)})
    return {"missing": sorted(found["all"] - reg_all),
            "stale": sorted(reg_all - found["all"]),
            "misclassified": sorted(misclassified)}


def classify_with_registry(found: dict,
                           assignments: dict[str, str] | None = None) -> dict:
    """--write 用：把 standalone_other 按现行注册表归属分到 文本段/标量；
    注册表也没有的新键由 assignments（--assign KEY=scalar|dynamic|text，
    人工决策）指定归属；仍未指定的进 unknown 拒绝落盘。"""
    from .report_registry import SCALAR_PLACEHOLDERS, TEXT_SECTION_PLACEHOLDERS
    other = set(found["standalone_other"])
    scalars = found["scalars"] | (other & SCALAR_PLACEHOLDERS)
    dynamic = set(found["dynamic_tables"])
    text = other & TEXT_SECTION_PLACEHOLDERS
    resolved = scalars | dynamic | text
    unknown = sorted(other - resolved)
    out_unknown: list[str] = []
    for k in unknown:
        cls = (assignments or {}).get(k)
        if cls == "scalar":
            scalars.add(k)
        elif cls == "dynamic":
            dynamic.add(k)
        elif cls == "text":
            text.add(k)
        else:
            out_unknown.append(k)
    return {"scalars": scalars, "dynamic_tables": dynamic,
            "text_sections": text, "unknown": out_unknown}
